from openpyxl import load_workbook
import datetime
import math
import os
import tempfile

# FOR logging costs when i was in london 2022-23

# load and activate the sheet
path = r"C:\Users\tamme\onedrive\desktop\files\docs\Money.xlsx"
wb = load_workbook(path)
ws = wb.active
# wb.active is whichever tab Excel had selected when it last saved, so say where
# the row is going rather than quietly writing to the wrong sheet
print(f"Sheet: {ws.title}")

# get the info to add
date = datetime.datetime.now().strftime("%m/%d/%Y")


def ask_amount():
    """Ask until the amount parses, so a typo doesn't discard the whole entry."""
    while True:
        raw = input("Enter amount: ").strip().lstrip("$").replace(",", "")
        try:
            value = float(raw)
        except ValueError:
            print("  not a number, try again")
            continue
        if not math.isfinite(value):
            # openpyxl writes nan/inf as an empty numeric cell
            print("  needs to be a finite amount, try again")
            continue
        return value


amount = ask_amount()
reason, details = map(input, ["Enter reason: ", "Enter details: "])

# add the info to a new row at the bottom. ws.append() starts after ws.max_row,
# which counts rows that carry only formatting -- so find the last row that
# actually holds a value.
last = 0
for cells in ws.iter_rows():
    if any(cell.value is not None for cell in cells):
        last = cells[0].row
target = last + 1
for column, value in enumerate([date, amount, reason, details], start=1):
    ws.cell(row=target, column=column, value=value)

# a row appended past a Table's range is invisible to =SUM(Table[Amount]) and to
# any pivot over it, so grow the range to take the new row in
for table in ws.tables.values():
    first, end = table.ref.split(":")
    column = "".join(ch for ch in end if ch.isalpha())
    if int("".join(ch for ch in end if ch.isdigit())) == target - 1:
        table.ref = f"{first}:{column}{target}"

# save. openpyxl truncates the destination before it starts serializing, so an
# interruption part way through leaves no ledger at all -- write alongside it and
# swap in the new file only once it is complete.
folder = os.path.dirname(os.path.abspath(path))
handle, temp = tempfile.mkstemp(dir=folder, prefix=".Money-", suffix=".xlsx")
os.close(handle)
try:
    wb.save(temp)
    os.replace(temp, path)
except BaseException:
    if os.path.exists(temp):
        os.remove(temp)
    raise

print(f"Logged {amount} ({reason}) to {ws.title} row {target}")
